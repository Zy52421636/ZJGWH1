import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import io

# 页面配置
st.set_page_config(page_title="气瓶收发统计系统", layout="centered")
st.title("📊 气瓶收发自动化统计系统")

# 上传组件
uploaded_file = st.file_uploader("请上传您的月报表文件 (.xlsx 或 .xlsm)", type=["xlsx", "xlsm"])

def categorize(c_type, p_desc):
    """核心分类逻辑，与之前的要求完全一致"""
    c_type = str(c_type).strip().upper()
    p_desc = str(p_desc).strip().upper()
    
    if c_type == "CYLINDER":
        return "Cylinder"
    elif c_type == "TONTANK":
        if "NH3" in p_desc and "930" in pdesc:
            return "TT 930 L"
        else:
            return "TT 440L"
    elif c_type == "BUNDLE":
        if "SIH4" in p_desc and "355KG" in pdesc:
            return "Bundle 28 cyl"
        else:
            return "Bundle 16 cyl"
    elif c_type == "NOT AVAILABLE":
        if "16*50" in p_desc:
            return "Bundle 28 cyl"
        elif "440L以下" in pdesc:
            return "Cylinder"
        elif "440L以上" in pdesc:
            if "NH3" in pdesc and "930" in pdesc:
                return "TT 930 L"
            else:
                return "TT 440L"
    return None

if uploaded_file is not None:
    st.info("正在处理数据，请稍候...")
    
    try:
        # 1. 使用 Pandas 读取满瓶和空瓶数据
        df_full = pd.read_excel(uploaded_file, sheet_name="满瓶入库")
        df_empty = pd.read_excel(uploaded_file, sheet_name="空瓶入库")
        
        # 2. 应用分类规则
        df_full['Category'] = df_full.apply(lambda row: categorize(row.get('CYLINDER_TYPE_DESCR', ''), row.get('PROD_DESCR', '')), axis=1)
        df_empty['Category'] = df_empty.apply(lambda row: categorize(row.get('CYLINDER_TYPE_DESCR', ''), row.get('PROD_DESCR', '')), axis=1)
        
        # 3. 统计数量
        counts_full = df_full['Category'].value_counts().to_dict()
        counts_empty = df_empty['Category'].value_counts().to_dict()
        
        # 在网页前端展示统计结果
        col1, col2 = st.columns(2)
        with col1:
            st.subheader("🟢 收满瓶统计")
            st.dataframe(pd.DataFrame(list(counts_full.items()), columns=['类别', '数量']))
        with col2:
            st.subheader("⚪ 收空瓶统计")
            st.dataframe(pd.DataFrame(list(counts_empty.items()), columns=['类别', '数量']))
        
        # 4. 将结果写回到原 Excel 文件的“数据统计”Sheet中
        # keep_vba=True 确保原有的宏代码不会丢失
        wb = load_workbook(uploaded_file, keep_vba=True)
        ws_stat = wb["数据统计"]
        
        def update_stats(ws, section_name, counts_dict):
            # 寻找大标题行
            start_row = None
            for r in range(1, ws.max_row + 1):
                cell_val = str(ws.cell(row=r, column=1).value or "")
                if section_name in cell_val:
                    start_row = r
                    break
            
            if start_row:
                cat_col = 2 # B列是类别
                qty_col = 3 # C列是瓶数
                
                for r in range(start_row + 1, start_row + 15):
                    a_val = str(ws.cell(row=r, column=1).value or "").strip()
                    row_label = str(ws.cell(row=r, column=cat_col).value or "").strip()
                    
                    # 遇到下一个A列大标题则停止
                    if a_val != "":
                        break
                        
                    # 精准匹配并填入数量 (忽略大小写和空格)
                    for k, v in counts_dict.items():
                        if row_label.replace(" ", "").upper() == str(k).replace(" ", "").upper():
                            ws.cell(row=r, column=qty_col).value = v
                            break

        update_stats(ws_stat, "收满瓶", counts_full)
        update_stats(ws_stat, "收空瓶", counts_empty)
        
        # 5. 生成供用户下载的最终文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        st.success("数据统计并写入完成！请点击下方按钮下载最终报表。")
        
        # 保持原文件后缀名
        file_ext = uploaded_file.name.split('.')[-1]
        st.download_button(
            label="📥 下载处理后的月报表",
            data=output,
            file_name=f"已统计_{uploaded_file.name}",
            mime="application/vnd.ms-excel.sheet.macroEnabled.12" if file_ext == 'xlsm' else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    except Exception as e:
        st.error(f"处理过程中出现错误：{e}")
        st.info("提示：请确认上传的表格中包含【满瓶入库】、【空瓶入库】和【数据统计】Sheet。")
